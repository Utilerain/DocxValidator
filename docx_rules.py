from dataclasses import dataclass
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from enum import Enum
from typing import List, Optional
from tqdm import tqdm
import re
import math

class Severity(Enum):
    ERROR = "ERR"
    WARNING = "WARN"

@dataclass
class Violation:
    rule_name: str
    severity: Severity
    location: str
    message: str

class BaseRule:
    def check(self, doc: Document) -> List[Violation]:
        raise NotImplementedError

def is_toc_line(text: str) -> bool:
    """
    Проверяет, является ли строка элементом оглавления
    (номер страницы в конце отделяется табуляцией или точками-заполнителями).
    """
    return bool(re.search(r'(\t|\.{2,}|…)\s*\d+\s*$', text))

def classify_paragraphs(doc: Document) -> List:
    """
    Размечает абзацы по типам, чтобы не натравливать MainTextFormatRule
    и другие правила на титульный лист и строки Содержания.

    Возвращает список размеченных абзацев: python-docx пересоздаёт обёртки
    Paragraph при каждом обращении к doc.paragraphs, поэтому флаги,
    выставленные на обёртках из другого списка, теряются.
    """
    classified = list(doc.paragraphs)
    in_title_page = True
    in_toc = False

    for p in classified:
        text = p.text.strip()

        # Определение начала оглавления (конец титульного листа)
        if "СОДЕРЖАНИЕ" in text.upper() or "ОГЛАВЛЕНИЕ" in text.upper():
            in_title_page = False
            in_toc = True
            p.is_title_page = False
            p.is_toc = True
            p.is_toc_heading = True
            continue

        # Определение конца Содержания (появление первого реального раздела)
        if in_toc and any(kw in text.upper() for kw in ["ВВЕДЕНИЕ", "ГЛАВА", "ЗАКЛЮЧЕНИЕ"]) and not is_toc_line(text):
            in_toc = False

        p.is_title_page = in_title_page
        p.is_toc = in_toc or is_toc_line(text)

    return classified

class PageMarginsRule(BaseRule):
    def __init__(self, top=2.0, bottom=2.0, left=3.0, right=1.5):
        self.top = top
        self.bottom = bottom
        self.left = left
        self.right = right

    def check(self, doc: Document) -> List[Violation]:
        violations = []
        for i, section in enumerate(doc.sections, 1):
            top_cm = round(section.top_margin.cm, 2) if section.top_margin else None
            bottom_cm = round(section.bottom_margin.cm, 2) if section.bottom_margin else None
            left_cm = round(section.left_margin.cm, 2) if section.left_margin else None
            right_cm = round(section.right_margin.cm, 2) if section.right_margin else None

            if top_cm != self.top or bottom_cm != self.bottom or left_cm != self.left or right_cm != self.right:
                violations.append(Violation(
                    rule_name="Поля страницы",
                    severity=Severity.ERROR,
                    location="Общее",
                    message=f"Ожидалось: Л: {self.left}cм, П: {self.right}cм, В: {self.top}cм, Н: {self.bottom}см.\n" +
                            f"Фактически: Л:{left_cm}см, П:{right_cm}см, В:{top_cm}см, Н:{bottom_cm}см."
                ))
        return violations

class MainTextFormatRule(BaseRule):
    def __init__(
        self,
        target_font_name: str = "Times New Roman",
        target_font_size: float = 14.0,
        target_line_spacing: float = 1.5,
        target_first_line_indent_cm: float = 1.25,
        target_space_before_pt: float = 0.0,
        target_space_after_pt: float = 0.0,
        allow_formatting_styles: bool = False,  # Разрешать ли bold/italic
        indent_tolerance_cm: float = 0.05       # Допустимая погрешность отступа
    ):
        self.font_name = target_font_name
        self.font_size = target_font_size
        self.line_spacing = target_line_spacing
        self.first_line_indent_cm = target_first_line_indent_cm
        self.space_before_pt = target_space_before_pt
        self.space_after_pt = target_space_after_pt
        self.allow_formatting_styles = allow_formatting_styles
        self.indent_tolerance_cm = indent_tolerance_cm

    def _get_effective_font_name(self, run, paragraph) -> Optional[str]:
        if run.font.name:
            return run.font.name
        if paragraph.style and paragraph.style.font and paragraph.style.font.name:
            return paragraph.style.font.name
        return None

    def _get_effective_font_size(self, run, paragraph) -> Optional[float]:
        if run.font.size:
            return run.font.size.pt
        if paragraph.style and paragraph.style.font and paragraph.style.font.size:
            return paragraph.style.font.size.pt
        return None

    def _get_effective_alignment(self, paragraph):
        if paragraph.alignment is not None:
            return paragraph.alignment
        if paragraph.style and paragraph.style.paragraph_format.alignment is not None:
            return paragraph.style.paragraph_format.alignment
        return None

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "MainTextFormatRule"

        paragraphs = classify_paragraphs(doc)

        for p_idx, paragraph in enumerate(paragraphs, start=1):
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue  # Пропускаем титульный лист и оглавление

            text = paragraph.text.strip()
            if not text:
                continue

            loc = f"Абзац {p_idx} ('{text[:30]}...')"
            fmt = paragraph.paragraph_format

            # 1. Выравнивание (по ширине / JUSTIFY)
            alignment = self._get_effective_alignment(paragraph)
            if alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Выравнивание текста должно быть по ширине (JUSTIFY)."
                    )
                )

            # 2. Красная строка (1,25 см)
            indent = fmt.first_line_indent
            if indent is None and paragraph.style and paragraph.style.paragraph_format.first_line_indent:
                indent = paragraph.style.paragraph_format.first_line_indent

            if indent is None:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message=f"Не задана красная строка (ожидается {self.first_line_indent_cm} см)."
                    )
                )
            else:
                indent_cm = indent.cm
                if abs(indent_cm - self.first_line_indent_cm) > self.indent_tolerance_cm:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message=f"Отступ красной строки равен {round(indent_cm, 2)} см, а должен быть {self.first_line_indent_cm} см."
                        )
                    )

            # 3. Междустрочный интервал (1.5)
            line_sp = fmt.line_spacing
            if line_sp is None and paragraph.style and paragraph.style.paragraph_format.line_spacing:
                line_sp = paragraph.style.paragraph_format.line_spacing

            if line_sp is not None and abs(float(line_sp) - self.line_spacing) > 0.05:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message=f"Междустрочный интервал {line_sp}, ожидается {self.line_spacing}."
                    )
                )

            # 4. Отступы до и после абзаца (0 pt)
            space_before = fmt.space_before.pt if fmt.space_before else 0.0
            space_after = fmt.space_after.pt if fmt.space_after else 0.0

            if abs(space_before - self.space_before_pt) > 0.1:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"Отступ перед абзацем равен {space_before} pt (ожидается {self.space_before_pt} pt)."
                    )
                )

            if abs(space_after - self.space_after_pt) > 0.1:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"Отступ после абзаца равен {space_after} pt (ожидается {self.space_after_pt} pt)."
                    )
                )

            # 5. Запрет висячих строк
            if fmt.widow_control is False:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="Отключен запрет висячих строк (widow_control)."
                    )
                )

            # 6. Проверка свойств текста (Runs): шрифт, кегль, начертания
            for r_idx, run in enumerate(paragraph.runs, start=1):
                run_text = run.text.strip()
                if not run_text:
                    continue

                run_loc = f"{loc}, фрагмент {r_idx} ('{run_text[:15]}')"

                font_name = self._get_effective_font_name(run, paragraph)
                if font_name and font_name != self.font_name:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=run_loc,
                            message=f"Шрифт '{font_name}', а должен быть '{self.font_name}'."
                        )
                    )

                font_size = self._get_effective_font_size(run, paragraph)
                if font_size and abs(font_size - self.font_size) > 0.1:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=run_loc,
                            message=f"Размер шрифта {font_size} pt, а должен быть {self.font_size} pt."
                        )
                    )

                if not self.allow_formatting_styles:
                    if run.bold:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=run_loc,
                                message="В тексте использовано полужирное начертание (bold)."
                            )
                        )
                    if run.italic:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=run_loc,
                                message="В тексте использован курсив (italic)."
                            )
                        )
                    if run.underline:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=run_loc,
                                message="В тексте использовано подчеркивание (underline)."
                            )
                        )

        return violations

class PageNumberingRule(BaseRule):
    def __init__(
        self,
        target_font_name: str = "Times New Roman",
        target_font_size: float = 10.0,
        target_alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ):
        self.font_name = target_font_name
        self.font_size = target_font_size
        self.alignment = target_alignment

    def _has_page_number_field(self, header_or_paragraph) -> bool:
        xml_str = header_or_paragraph._element.xml
        return ("instrText" in xml_str and "PAGE" in xml_str) or ("w:fldSimple" in xml_str and "PAGE" in xml_str)

    def _get_effective_font_name(self, run, paragraph) -> Optional[str]:
        if run.font.name:
            return run.font.name
        if paragraph.style and paragraph.style.font and paragraph.style.font.name:
            return paragraph.style.font.name
        return None

    def _get_effective_font_size(self, run, paragraph) -> Optional[float]:
        if run.font.size:
            return run.font.size.pt
        if paragraph.style and paragraph.style.font and paragraph.style.font.size:
            return paragraph.style.font.size.pt
        return None

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "PageNumberingRule"

        if not doc.sections:
            return violations

        first_section = doc.sections[0]

        if first_section.different_first_page_header_footer:
            if self._has_page_number_field(first_section.first_page_header):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location="Секция 1 (Колонтитул первой страницы)",
                        message="На титульном листе присутствует нумерация страниц."
                    )
                )
        else:
            if self._has_page_number_field(first_section.header):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location="Секция 1 (Верхний колонтитул)",
                        message="Не включена опция 'Особый колонтитул для первой страницы'. На титульном листе не должно быть нумерации."
                    )
                )

        for sec_idx, section in enumerate(doc.sections, start=1):
            header = section.header

            if not self._has_page_number_field(header):
                continue

            paragraphs = header.paragraphs

            if len(paragraphs) > 1:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=f"Секция {sec_idx} (Верхний колонтитул)",
                        message=f"В верхнем колонтитуле с нумерацией обнаружены лишние пустые строки (всего строк: {len(paragraphs)}, должна быть 1)."
                    )
                )

            for p_idx, paragraph in enumerate(paragraphs, start=1):
                if "\n" in paragraph.text:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Секция {sec_idx}, колонтитул, абзац {p_idx}",
                            message="В строке нумерации присутствуют лишние переносы строк."
                        )
                    )

                alignment = paragraph.alignment
                if alignment is None and paragraph.style and paragraph.style.paragraph_format:
                    alignment = paragraph.style.paragraph_format.alignment

                if alignment != self.alignment:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Секция {sec_idx}, колонтитул, абзац {p_idx}",
                            message="Нумерация страниц должна быть выровнена по правому краю (правый верхний угол)."
                        )
                    )

                for r_idx, run in enumerate(paragraph.runs, start=1):
                    font_name = self._get_effective_font_name(run, paragraph)
                    if font_name and font_name != self.font_name:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.ERROR,
                                location=f"Секция {sec_idx}, колонтитул, абзац {p_idx}, фрагмент {r_idx}",
                                message=f"Шрифт нумерации '{font_name}', а должен быть '{self.font_name}'."
                            )
                        )

                    font_size = self._get_effective_font_size(run, paragraph)
                    if font_size and abs(font_size - self.font_size) > 0.1:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.ERROR,
                                location=f"Секция {sec_idx}, колонтитул, абзац {p_idx}, фрагмент {r_idx}",
                                message=f"Размер шрифта нумерации {font_size} pt, а должен быть {self.font_size} pt."
                            )
                        )

        return violations

class TitlePageRule(BaseRule):
    def _get_title_page_paragraphs(self, doc: Document) -> List:
        title_paragraphs = []
        for p in doc.paragraphs:
            text_upper = p.text.strip().upper()

            has_page_break = any(
                'w:br' in run._element.xml and 'w:type="page"' in run._element.xml
                for run in p.runs
            )

            if text_upper in ["СОДЕРЖАНИЕ", "ВВЕДЕНИЕ", "ОГЛАВЛЕНИЕ"] or has_page_break:
                break

            title_paragraphs.append(p)

            if re.search(r'\b20\d{2}\b', text_upper) and len(text_upper) < 30:
                break

        return title_paragraphs

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "TitlePageRule"

        title_paragraphs = self._get_title_page_paragraphs(doc)
        total_p = len(title_paragraphs)

        if total_p < 5:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location="Титульный лист",
                    message="Титульный лист не найден или содержит слишком мало текста."
                )
            )
            return violations

        topic_idx: Optional[int] = None
        topic_paragraph = None
        has_institution_header = False
        has_work_type = False
        has_year = False

        for idx, p in enumerate(title_paragraphs):
            text = p.text.strip()
            text_upper = text.upper()

            if not text:
                continue

            align = p.alignment
            if align is None and p.style and p.style.paragraph_format:
                align = p.style.paragraph_format.alignment

            if any(word in text_upper for word in ["МИНИСТЕРСТВО", "ОБРАЗОВАТЕЛЬНОЕ", "КОЛЛЕДЖ", "УНИВЕРСИТЕТ", "ИНСТИТУТ"]):
                has_institution_header = True
                if align != WD_ALIGN_PARAGRAPH.CENTER:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Титульный лист, абзац {idx+1} ('{text[:30]}...')",
                            message="Шапка учебного заведения должна быть выровнена по центру."
                        )
                    )

            if "ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА" in text_upper or "ДИПЛОМНЫЙ ПРОЕКТ" in text_upper:
                has_work_type = True

            is_bold = any(run.bold for run in p.runs if run.text.strip())
            is_topic_candidate = (
                ("РАЗРАБОТКА" in text_upper or "ТЕМА" in text_upper or "ПРОЕКТИРОВАНИЕ" in text_upper or is_bold)
                and not has_work_type
                and idx > 2
            )

            if is_topic_candidate and topic_paragraph is None and "МИНИСТЕРСТВО" not in text_upper and "КОЛЛЕДЖ" not in text_upper:
                topic_idx = idx
                topic_paragraph = p

            if re.search(r'\b20\d{2}\b', text_upper):
                has_year = True
                if align != WD_ALIGN_PARAGRAPH.CENTER:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Титульный лист, абзац {idx+1} ('{text}')",
                            message="Город и год внизу титульного листа должны быть выровнены по центру."
                        )
                    )

            if p.paragraph_format.first_line_indent and p.paragraph_format.first_line_indent.cm > 0.1:
                if align == WD_ALIGN_PARAGRAPH.CENTER or "МИНИСТЕРСТВО" in text_upper or "РАЗРАБОТКА" in text_upper:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"Титульный лист, абзац {idx+1} ('{text[:30]}...')",
                            message="У центрального блока титульного листа установлена красная строка. Значение должно быть 0."
                        )
                    )

        if topic_paragraph is not None and topic_idx is not None:
            topic_align = topic_paragraph.alignment
            if topic_align is None and topic_paragraph.style and topic_paragraph.style.paragraph_format:
                topic_align = topic_paragraph.style.paragraph_format.alignment

            if topic_align != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=f"Титульный лист, абзац {topic_idx+1} ('{topic_paragraph.text[:30]}...')",
                        message="Тема дипломного проекта должна быть выровнена по центру."
                    )
                )

            topic_has_bold = any(run.bold for run in topic_paragraph.runs if run.text.strip())
            if not topic_has_bold:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=f"Титульный лист, абзац {topic_idx+1}",
                        message="Тема дипломного проекта должна быть выделена полужирным шрифтом."
                    )
                )

            relative_position = topic_idx / total_p
            if relative_position < 0.20 or relative_position > 0.75:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=f"Титульный лист, абзац {topic_idx+1}",
                        message=f"Тема дипломного проекта расположена слишком {'высоко' if relative_position < 0.20 else 'низко'}. Она должна быть примерно в центре титульного листа."
                    )
                )
        else:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location="Титульный лист",
                    message="Тема дипломного проекта не найдена на титульном листе."
                )
            )

        if not has_institution_header:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location="Титульный лист",
                    message="Не найдена шапка учебного заведения / министерства."
                )
            )

        if not has_year:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location="Титульный лист",
                    message="Внизу титульного листа не найден город и год."
                )
            )

        return violations

class TableOfContentsRule(BaseRule):
    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "TableOfContentsRule"

        toc_header_index: Optional[int] = None

        for idx, p in enumerate(doc.paragraphs):
            clean_text = p.text.strip().upper()
            if clean_text in ["СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ"]:
                toc_header_index = idx
                loc = f"Заголовок оглавления (абзац {idx+1})"

                align = p.alignment
                if align is None and p.style and p.style.paragraph_format:
                    align = p.style.paragraph_format.alignment
                if align != WD_ALIGN_PARAGRAPH.CENTER:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message="Заголовок 'Содержание' должен быть выровнен по центру."
                        )
                    )

                is_bold = any(run.bold for run in p.runs if run.text.strip())
                if not is_bold and p.style and p.style.font:
                    is_bold = p.style.font.bold or False
                if not is_bold:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message="Название 'Содержание' должно быть выделено полужирным шрифтом."
                        )
                    )

                fmt = p.paragraph_format
                first_indent = fmt.first_line_indent
                if first_indent is None and p.style and p.style.paragraph_format:
                    first_indent = p.style.paragraph_format.first_line_indent
                if first_indent is None or first_indent.cm <= 0.05:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message="У заголовка 'Содержание' должна быть установлена красная строка."
                        )
                    )
                break

        if toc_header_index is None:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location="Документ",
                    message="Раздел 'Содержание' (или 'Оглавление') не найден в документе."
                )
            )
            return violations

        toc_entries = []
        for p_idx in range(toc_header_index + 1, len(doc.paragraphs)):
            p = doc.paragraphs[p_idx]
            text = p.text.strip()
            if not text:
                continue

            is_toc_field = 'TOC' in p._element.xml or 'w:fldSimple' in p._element.xml or 'w:instrText' in p._element.xml
            has_page_num_at_end = re.search(r'\d+\s*$', text) is not None

            if not has_page_num_at_end and not is_toc_field and text.upper() in ["ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "ПРИЛОЖЕНИЕ"]:
                break

            has_page_break = any(
                'w:br' in run._element.xml and 'w:type="page"' in run._element.xml
                for run in p.runs
            )

            toc_entries.append((p_idx + 1, p))

            if has_page_break:
                break

        for p_num, p in toc_entries:
            text = p.text.strip()
            loc = f"Оглавление, абзац {p_num} ('{text[:30]}...')"

            if re.search(r'\.{3,}', text) or re.search(r'(\.\s*){4,}', text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Проставление точек не допустимо! Используйте стандартную табуляцию с заполнителем или авто-оглавление Word."
                    )
                )

            fmt = p.paragraph_format
            align = p.alignment
            if align is None and p.style and p.style.paragraph_format:
                align = p.style.paragraph_format.alignment

            if align != WD_ALIGN_PARAGRAPH.JUSTIFY:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Пункт оглавления должен быть выровнен по ширине (JUSTIFY)."
                    )
                )

            first_indent = fmt.first_line_indent
            if first_indent is None and p.style and p.style.paragraph_format:
                first_indent = p.style.paragraph_format.first_line_indent

            if first_indent is not None and first_indent.cm > 0.05:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message=f"У пункта оглавления не должно быть красной строки (сейчас {round(first_indent.cm, 2)} см)."
                    )
                )

            num_match = re.match(r'^(Глава\s+\d+\.|Глава\s+\d+|\d+(\.\d+)*\.)(\s+)', text, re.IGNORECASE)
            if num_match:
                spaces = num_match.group(3)
                if len(spaces) > 1:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message="Между номером главы/параграфа и названием должен быть строго один пробел."
                        )
                    )

        return violations

class SectionFormatRule(BaseRule):
    MAJOR_HEADER_REGEX = r'^(ГЛАВА\s+\d+|СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ)|ПРИЛОЖЕНИЕ)'
    SUBHEADING_REGEX = r'^\d+(\.\d+)+\.?\s+'

    def _is_empty(self, paragraph) -> bool:
        return not paragraph.text.strip()

    def _is_bold(self, paragraph) -> bool:
        runs_with_text = [r for r in paragraph.runs if r.text.strip()]
        if runs_with_text:
            return all(r.bold for r in runs_with_text)
        if paragraph.style and paragraph.style.font:
            return paragraph.style.font.bold or False
        return False

    def _is_centered(self, paragraph) -> bool:
        align = paragraph.alignment
        if align is None and paragraph.style and paragraph.style.paragraph_format:
            align = paragraph.style.paragraph_format.alignment
        return align == WD_ALIGN_PARAGRAPH.CENTER

    def _has_red_line(self, paragraph) -> bool:
        fmt = paragraph.paragraph_format
        indent = fmt.first_line_indent
        if indent is None and paragraph.style and paragraph.style.paragraph_format:
            indent = paragraph.style.paragraph_format.first_line_indent
        return indent is not None and indent.cm > 0.05

    def _has_page_break_before(self, paragraphs: List, idx: int) -> bool:
        if idx == 0:
            return True

        p = paragraphs[idx]

        if p.paragraph_format.page_break_before:
            return True

        prev_p = paragraphs[idx - 1]
        xml_str = prev_p._element.xml
        if 'w:br' in xml_str and 'w:type="page"' in xml_str:
            return True

        if 'w:sectPr' in xml_str:
            return True

        return False

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "SectionFormatRule"

        paragraphs = classify_paragraphs(doc)

        for idx, p in enumerate(paragraphs):
            text = p.text.strip()
            if not text:
                continue

            # Игнорируем титульный лист и строки оглавления
            if getattr(p, 'is_title_page', False) or getattr(p, 'is_toc', False) or is_toc_line(text):
                continue

            text_upper = text.upper()
            is_major = bool(re.match(self.MAJOR_HEADER_REGEX, text_upper, re.IGNORECASE))
            is_sub = bool(re.match(self.SUBHEADING_REGEX, text, re.IGNORECASE))

            if not (is_major or is_sub):
                continue

            header_type = "Глава/Крупный раздел" if is_major else "Параграф"
            loc = f"{header_type} (абзац {idx+1}: '{text[:30]}...')"

            # 1. Проверка полужирного начертания
            if not self._is_bold(p):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message=f"{header_type} должен быть выделен полужирным шрифтом."
                    )
                )

            # 2. Проверка выравнивания по центру
            if not self._is_centered(p):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message=f"Заголовок ({header_type.lower()}) должен быть выровнен по центру."
                    )
                )

            # 3. Проверка красной строки (не должно быть для выравнивания по центру)
            if self._is_centered(p):
                if self._has_red_line(p):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message=f"У центрированного заголовка ({header_type.lower()}) не должно быть красной строки (абзацного отступа)."
                        )
                    )
            else:
                if not self._has_red_line(p):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message=f"У заголовка ({header_type.lower()}) должна быть установлена красная строка."
                        )
                    )

            # 4. Проверка правила «с новой страницы»
            starts_new_page = self._has_page_break_before(paragraphs, idx)
            if is_major and not starts_new_page:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Крупные разделы (Глава, Содержание, Введение, Заключение, Список источников) должны начинаться с новой страницы."
                    )
                )
            elif is_sub and p.paragraph_format.page_break_before:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="Параграф должен продолжаться с места окончания предыдущего (не должен принудительно начинаться с новой страницы)."
                    )
                )

            # 5. Проверка Enter'ов перед заголовком
            if is_sub:
                prev_non_empty_idx = idx - 1
                while prev_non_empty_idx >= 0 and self._is_empty(paragraphs[prev_non_empty_idx]):
                    prev_non_empty_idx -= 1

                prev_p_text = paragraphs[prev_non_empty_idx].text.strip().upper() if prev_non_empty_idx >= 0 else ""
                after_chapter = bool(re.match(self.MAJOR_HEADER_REGEX, prev_p_text, re.IGNORECASE))

                empty_count_before = 0
                check_i = idx - 1
                while check_i >= 0 and self._is_empty(paragraphs[check_i]):
                    empty_count_before += 1
                    check_i -= 1

                if after_chapter:
                    if empty_count_before != 1:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=loc,
                                message=f"После названия главы перед параграфом должен быть ровно 1 Enter (обнаружено пустых строк: {empty_count_before})."
                            )
                        )
                else:
                    if empty_count_before != 2 and not starts_new_page:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=loc,
                                message=f"Перед началом параграфа должно ставиться 2 Enter'а (обнаружено пустых строк: {empty_count_before})."
                            )
                        )

            # 6. Проверка Enter'а после заголовка
            empty_count_after = 0
            check_after = idx + 1
            while check_after < len(paragraphs) and self._is_empty(paragraphs[check_after]):
                empty_count_after += 1
                check_after += 1

            if empty_count_after != 1:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"После названия {header_type.lower()} должен оставаться ровно 1 Enter (обнаружено пустых строк: {empty_count_after})."
                    )
                )

        return violations

class SectionPageFillRule(BaseRule):
    MAJOR_HEADER_REGEX = r'^(ГЛАВА\s+\d+|СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ)|ПРИЛОЖЕНИЕ)'

    def __init__(
        self,
        min_fill_ratio: float = 0.70,
        chars_per_line: int = 70,
        lines_per_full_page: int = 28
    ):
        self.min_fill_ratio = min_fill_ratio
        self.chars_per_line = chars_per_line
        self.lines_per_full_page = lines_per_full_page
        self.min_lines_required = math.ceil(lines_per_full_page * min_fill_ratio)

    def _has_page_break(self, paragraph) -> bool:
        if paragraph.paragraph_format.page_break_before:
            return True
        xml_str = paragraph._element.xml
        return 'w:br' in xml_str and 'w:type="page"' in xml_str

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "SectionPageFillRule"

        paragraphs = classify_paragraphs(doc)

        section_blocks = []
        current_block = []

        for idx, p in enumerate(paragraphs):
            text = p.text.strip()
            text_upper = text.upper()

            # Игнорируем титульный лист и оглавление
            if getattr(p, 'is_title_page', False) or getattr(p, 'is_toc', False) or is_toc_line(text):
                continue

            is_major = bool(re.match(self.MAJOR_HEADER_REGEX, text_upper, re.IGNORECASE))
            has_break = self._has_page_break(p)

            if (is_major or has_break) and current_block:
                section_blocks.append((idx, current_block))
                current_block = []

            if text or 'w:drawing' in p._element.xml or 'w:pict' in p._element.xml:
                current_block.append(p)

        for end_idx, block_paragraphs in section_blocks:
            if not block_paragraphs:
                continue

            total_estimated_lines = 0

            for p in block_paragraphs:
                p_text = p.text.strip()

                if p_text:
                    lines_in_p = max(1, math.ceil(len(p_text) / self.chars_per_line))
                    total_estimated_lines += lines_in_p

                xml_str = p._element.xml
                if 'w:drawing' in xml_str or 'w:pict' in xml_str:
                    total_estimated_lines += 10

            fill_ratio = min(1.0, total_estimated_lines / self.lines_per_full_page)

            if fill_ratio < self.min_fill_ratio:
                last_text = block_paragraphs[-1].text.strip()[:30]
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=f"Конец раздела перед абзацем {end_idx} ('{last_text}...')",
                        message=(
                            f"При окончании раздела страница заполнена ориентировочно на {int(fill_ratio * 100)}% "
                            f"(~{total_estimated_lines} из {self.lines_per_full_page} строк). "
                            f"Допустимо заполнение листа не менее, чем на {int(self.min_fill_ratio * 100)}%."
                        )
                    )
                )

        return violations

class ReferencesListRule(BaseRule):
    HEADER_REGEX = r'^(СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ))'
    NEXT_SECTION_REGEX = r'^(ПРИЛОЖЕНИЕ|ГЛАВА\s+\d+|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СОДЕРЖАНИЕ)'
    URL_REGEX = r'(https?://[^\s]+|www\.[^\s]+|\[?электронный\s+ресурс\]?)'
    ACCESS_DATE_REGEX = r'дата\s+обращения:?\s*\d{2}\.\d{2}\.\d{4}'
    PERIODICAL_KEYWORDS = ['ЖУРНАЛ', 'ВЕСТНИК', 'ИЗВЕСТИЯ', 'ПЕРИОДИЧЕСКОЕ ИЗДАНИЕ', 'ВЫП.']

    def _get_sort_key(self, text: str) -> str:
        cleaned = re.sub(r'^\d+[\.\)]\s*', '', text).strip()
        cleaned = re.sub(r'^\s*\[\d+\]\s*', '', cleaned)
        return cleaned.lower()

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "ReferencesListRule"

        ref_header_idx = None
        for idx, p in enumerate(doc.paragraphs):
            text = p.text.strip().upper()
            if re.match(self.HEADER_REGEX, text) and len(text) < 100:
                ref_header_idx = idx

        if ref_header_idx is None:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location="Документ",
                    message="Раздел 'Список источников и литературы' не найден в документе."
                )
            )
            return violations

        items = []
        for p_idx in range(ref_header_idx + 1, len(doc.paragraphs)):
            p = doc.paragraphs[p_idx]
            text = p.text.strip()
            if not text:
                continue

            text_upper = text.upper()
            if re.match(self.NEXT_SECTION_REGEX, text_upper) and p_idx > ref_header_idx + 1:
                break

            items.append((p_idx + 1, text))

        if not items:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location=f"Абзац {ref_header_idx + 1}",
                    message="Раздел 'Список источников и литературы' пуст."
                )
            )
            return violations

        printed_items = []
        electronic_items = []
        seen_electronic = False

        for p_num, text in items:
            loc = f"Список источников, абзац {p_num} ('{text[:30]}...')"

            if not re.match(r'^\d+[\.\)]\s*', text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Источник должен быть частью нумерованного списка (начинаться с '1. ')."
                    )
                )

            clean_text = re.sub(r'^\d+[\.\)]\s*', '', text).strip()
            is_electronic = bool(re.search(self.URL_REGEX, text, re.IGNORECASE))

            if is_electronic:
                seen_electronic = True
                electronic_items.append((p_num, text))

                if not re.search(self.ACCESS_DATE_REGEX, text, re.IGNORECASE):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message="Для электронного ресурса обязательно должна быть указана дата обращения (например: 'дата обращения: 06.06.2025')."
                        )
                    )
            else:
                if seen_electronic:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=loc,
                            message="Печатное издание находится после электронного ресурса. Электронные ресурсы прописываются строго ПОСЛЕ печатных изданий."
                        )
                    )
                printed_items.append((p_num, text))

            text_upper = text.upper()
            if any(kw in text_upper for kw in self.PERIODICAL_KEYWORDS) and '//' in text:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="Обнаружены признаки периодического издания (статья/журнал). Использование периодических изданий не допускается."
                    )
                )

            initials_first_match = re.match(r'^([А-ЯA-Z]\.\s*[А-ЯA-Z]\.|[А-ЯA-Z]\.)\s+([А-ЯA-Z][а-яa-z]+)', clean_text)
            if initials_first_match:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"Нарушен порядок записи автора ('{initials_first_match.group(0)}'). Инициалы должны идти ПОСЛЕ фамилии: Фамилия И.О. (например, 'Достоевский Ф.М.')."
                    )
                )

        for i in range(1, len(printed_items)):
            prev_pnum, prev_text = printed_items[i - 1]
            curr_pnum, curr_text = printed_items[i]

            if self._get_sort_key(curr_text) < self._get_sort_key(prev_text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=f"Список источников, элемент {curr_pnum}",
                        message=f"Нарушен алфавитный порядок печатных изданий! Элемент '{curr_text[:25]}...' должен располагаться раньше '{prev_text[:25]}...'."
                    )
                )

        for i in range(1, len(electronic_items)):
            prev_pnum, prev_text = electronic_items[i - 1]
            curr_pnum, curr_text = electronic_items[i]

            if self._get_sort_key(curr_text) < self._get_sort_key(prev_text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=f"Список источников, элемент {curr_pnum}",
                        message=f"Нарушен алфавитный порядок электронных ресурсов! Элемент '{curr_text[:25]}...' должен располагаться раньше '{prev_text[:25]}...'."
                    )
                )

        return violations

class IntroductionContentRule(BaseRule):
    MAJOR_HEADER_REGEX = r'^(ГЛАВА\s+\d+|СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ)|ПРИЛОЖЕНИЕ)'

    # Обязательные элементы введения (слайд 203)
    REQUIRED_ELEMENTS = [
        ("Теоретическое описание", r'ТЕОРЕТИЧЕСК'),
        ("Актуальность", r'АКТУАЛЬНОСТ'),
        ("Тема", r'\bТЕМ[АЕУОЙ]\b'),
        ("Объект", r'\bОБЪЕКТ'),
        ("Предмет", r'\bПРЕДМЕТ'),
        ("Цель", r'\bЦЕЛЬ'),
        ("Задачи", r'\bЗАДАЧ'),
        ("Методы", r'\bМЕТОД'),
        ("Теоретическая и практическая значимость", r'ЗНАЧИМОСТ'),
        ("Структура работы", r'\bСТРУКТУР'),
    ]

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "IntroductionContentRule"

        classify_paragraphs(doc)

        intro_start = None
        for idx, p in enumerate(classify_paragraphs(doc)):
            text = p.text.strip()
            if getattr(p, 'is_title_page', False) or getattr(p, 'is_toc', False) or is_toc_line(text):
                continue
            if text.upper() == "ВВЕДЕНИЕ":
                intro_start = idx + 1
                break

        if intro_start is None:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location="Документ",
                    message="Раздел 'Введение' не найден в документе."
                )
            )
            return violations

        intro_parts = []
        for p in doc.paragraphs[intro_start:]:
            text = p.text.strip()
            text_upper = text.upper()

            # Конец введения: крупный заголовок либо любой центрированный абзац
            # (в некоторых работах разделы идут без слова 'ГЛАВА')
            if text and re.match(self.MAJOR_HEADER_REGEX, text_upper):
                break

            align = p.alignment
            if align is None and p.style and p.style.paragraph_format:
                align = p.style.paragraph_format.alignment
            if text and align == WD_ALIGN_PARAGRAPH.CENTER:
                break

            intro_parts.append(text)
        intro_text = " ".join(intro_parts).upper()

        for element_name, pattern in self.REQUIRED_ELEMENTS:
            if not re.search(pattern, intro_text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location="Введение",
                        message=f"Во введении не найден обязательный элемент: {element_name}."
                    )
                )

        return violations

class ParagraphConclusionRule(BaseRule):
    MAJOR_HEADER_REGEX = r'^(ГЛАВА\s+\d+|СОДЕРЖАНИЕ|ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ)|ПРИЛОЖЕНИЕ)'
    SUBHEADING_REGEX = r'^\d+(\.\d+)+\.?\s+'
    NUMBERED_SECTION_REGEX = r'^\d+\.?\s+'
    STOP_SECTION_REGEX = r'^(СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ)|ПРИЛОЖЕНИЕ)'
    SKIP_SECTION_REGEX = r'^(ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СОДЕРЖАНИЕ)'

    # Начало вывода по итогу параграфа / главы
    CONCLUSION_START_REGEX = (
        r'^(ТАКИМ ОБРАЗОМ|ИТАК\b|ПОДВОДЯ ИТОГ|ОБОБЩАЯ|РЕЗЮМИРУЯ|СЛЕДОВАТЕЛЬНО|'
        r'ПРОАНАЛИЗИРОВАВ|ВЫВОД[Ы]?\b|ВСЁ ВЫШЕСКАЗАННОЕ|ВЫШЕИЗЛОЖЕННОЕ)'
    )

    @staticmethod
    def iter_body_sequence(paragraphs: List, tables: List):
        """
        Последовательность элементов тела документа в порядке следования:
        ('p', абзац, индекс_абзаца) и ('tbl', таблица, None).
        Нужна, чтобы видеть таблицы между абзацами.
        """
        items = []
        p_idx = 0
        para_iter = iter(paragraphs)
        tbl_iter = iter(tables)
        body = None
        if paragraphs:
            body = paragraphs[0]._element.getparent()
        if body is None:
            return items
        for child in body.iterchildren():
            if child.tag == qn('w:p'):
                para = next(para_iter, None)
                if para is not None:
                    items.append(('p', para, p_idx))
                    p_idx += 1
            elif child.tag == qn('w:tbl'):
                table = next(tbl_iter, None)
                if table is not None:
                    items.append(('tbl', table, None))
        return items

    def _is_centered(self, paragraph) -> bool:
        align = paragraph.alignment
        if align is None and paragraph.style and paragraph.style.paragraph_format:
            align = paragraph.style.paragraph_format.alignment
        return align == WD_ALIGN_PARAGRAPH.CENTER

    def _header_type(self, paragraph) -> Optional[str]:
        text = paragraph.text.strip()
        if not text:
            return None
        text_upper = text.upper()
        if re.match(self.MAJOR_HEADER_REGEX, text_upper):
            return 'major'
        if re.match(self.SUBHEADING_REGEX, text):
            return 'sub'
        # Раздел без слова 'Глава' (например '1. Понятие ...'), оформленный по центру
        if re.match(self.NUMBERED_SECTION_REGEX, text) and self._is_centered(paragraph) and len(text) < 120:
            return 'major'
        return None

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "ParagraphConclusionRule"

        paragraphs = classify_paragraphs(doc)
        items = self.iter_body_sequence(paragraphs, doc.tables)

        starts = []
        for i, (kind, obj, p_idx) in enumerate(items):
            if kind != 'p':
                continue
            if getattr(obj, 'is_title_page', False) or getattr(obj, 'is_toc', False):
                continue
            header_type = self._header_type(obj)
            if header_type:
                starts.append((i, header_type, obj, p_idx))

        for pos, (start_idx, header_type, header_para, header_idx) in enumerate(starts):
            text_upper = header_para.text.strip().upper()

            if re.match(self.STOP_SECTION_REGEX, text_upper):
                break  # После списка источников и приложения выводы не проверяются
            if header_type != 'major' and header_type != 'sub':
                continue
            if re.match(self.SKIP_SECTION_REGEX, text_upper):
                continue  # Введение, заключение и содержание не являются параграфами/главами

            end_idx = starts[pos + 1][0] if pos + 1 < len(starts) else len(items)
            block = items[start_idx + 1:end_idx]

            last_para = None
            last_para_idx = None
            last_item_kind = None
            for kind, obj, p_idx in block:
                if kind == 'p':
                    if obj.text.strip():
                        last_para = obj
                        last_para_idx = p_idx
                        last_item_kind = 'p'
                else:
                    last_item_kind = 'tbl'

            if last_para is None:
                continue  # Блок без содержания (например, заголовок главы перед параграфом)

            header_text = header_para.text.strip()[:40]
            loc = f"Конец '{header_text}' (абзац {last_para_idx + 1})"

            # 1. Таблица в конце раздела
            if last_item_kind == 'tbl':
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="В конце параграфа не допустима таблица (она должна быть внутри текста раздела)."
                    )
                )

            # 2. Вывод по итогу раздела
            if not re.match(self.CONCLUSION_START_REGEX, last_para.text.strip().upper()):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="Параграф, глава всегда должны заканчиваться выводом (например: 'Таким образом, ...')."
                    )
                )

            # 3. Недопустимые элементы в конце раздела
            last_xml = last_para._element.xml
            if '<w:drawing' in last_xml or '<w:pict' in last_xml:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="В конце параграфа не допустимы рисунки."
                    )
                )

            if 'oMath' in last_xml:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="В конце параграфа не допустимы формулы."
                    )
                )

            if re.search(r'\[\d+', last_para.text):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="В конце параграфа не допустимы ссылки на источники."
                    )
                )

        return violations

class SourceReferencesRule(BaseRule):
    CHAPTER_REGEX = r'^ГЛАВА\s+\d+'
    STOP_SECTION_REGEX = r'^(СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ)|ПРИЛОЖЕНИЕ)'
    LINK_REGEX = re.compile(r'\[([^\[\]]{1,30})\]')
    VALID_LINK_REGEX = r'^\d+(,\s*с\.\s*\d+)?$'

    # Строки листингов кода не являются ссылками на источники
    CODE_LINE_REGEX = re.compile(
        r'^\s*('
        r'import\s|from\s+[A-Za-z_][\w.]*\s+import|'
        r'[A-Za-z_]\w*\s*(\[[^\]]*\])?\s*=[^=]|'
        r'print\s*\(|plt\.|def\s+[A-Za-z_]\w*\s*\(|return\b|'
        r'for\s+[A-Za-z_]\w*\s+in\s|while\s|elif\b|else:|#'
        r')'
    )

    CHARS_PER_LINE = 70
    LINES_PER_PAGE = 28

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "SourceReferencesRule"

        paragraphs = classify_paragraphs(doc)

        in_code = False
        for p_idx, paragraph in enumerate(paragraphs, start=1):
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue

            text = paragraph.text
            if not text.strip():
                continue

            if re.match(self.STOP_SECTION_REGEX, text.strip().upper()):
                break  # Список литературы и приложения не содержат ссылок-цитирований

            if self.CODE_LINE_REGEX.match(text):
                in_code = True
                continue

            if in_code:
                stripped = text.strip()
                # Конец листинга: абзац выглядит как обычный русский текст
                if re.search(r'[а-яё]{3}', stripped.lower()) and stripped.endswith('.') and ' ' in stripped:
                    in_code = False
                else:
                    continue

            for match in self.LINK_REGEX.finditer(text):
                content = match.group(1).strip()

                # Цитата состоит только из цифр, запятых, точек, пробелов и 'с.'
                if not re.search(r'\d', content) or not re.match(r'^[\d\s,с.]+$', content):
                    continue

                start, end = match.span()
                loc = f"Абзац {p_idx} ('{text.strip()[:30]}...')"

                if not re.match(self.VALID_LINK_REGEX, content):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message=f"Ссылка на источник '[{content}]' должна иметь формат [N] или [N, с.M] (например: [7, с.45])."
                        )
                    )
                    continue

                # Ссылка ставится в конце абзаца перед точкой
                tail = text[end:].strip()
                if tail not in ("", "."):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message=f"Ссылка на источник '[{content}]' должна стоять в конце абзаца перед точкой."
                        )
                    )

                # Перед ссылкой должен быть пробел
                if start > 0 and text[start - 1].isalnum():
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message=f"Перед ссылкой '[{content}]' должен быть пробел."
                        )
                    )

        # В теоретическом разделе ссылки должны присутствовать на каждой странице
        chapter_paragraphs = []
        in_chapter = False
        for p_idx, paragraph in enumerate(paragraphs, start=1):
            text = paragraph.text.strip()
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue

            if re.match(self.CHAPTER_REGEX, text.upper()):
                if in_chapter:
                    break  # Конец первой главы (теоретического раздела)
                in_chapter = True
                estimated_lines = 0.0
                continue

            if in_chapter:
                if re.match(r'^(ГЛАВА\s+\d+|ЗАКЛЮЧЕНИЕ|СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ)|ПРИЛОЖЕНИЕ)', text.upper()):
                    break
                chapter_paragraphs.append((p_idx, paragraph, estimated_lines))
                if text:
                    estimated_lines += math.ceil(len(text) / self.CHARS_PER_LINE)
                if '<w:drawing' in paragraph._element.xml or '<w:pict' in paragraph._element.xml:
                    estimated_lines += 10

        if chapter_paragraphs:
            pages_refs = {}
            for p_idx, paragraph, est_lines in chapter_paragraphs:
                page_no = int(est_lines // self.LINES_PER_PAGE)
                pages_refs.setdefault(page_no, 0)
                if re.search(r'\[\d+', paragraph.text):
                    pages_refs[page_no] += 1

            for page_no in sorted(pages_refs):
                if pages_refs[page_no] == 0:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"Теоретический раздел (Глава 1), ~страница {page_no + 1} от начала главы",
                            message="В теоретическом разделе ссылки на источники должны присутствовать на каждой странице."
                        )
                    )

        return violations

class AppendixFormatRule(BaseRule):
    HEADER_REGEX = r'^ПРИЛОЖЕНИЕ'

    @staticmethod
    def _get_paragraph_section_indices(doc: Document) -> List[int]:
        """Индекс секции для каждого абзаца doc.paragraphs (секции разделяются sectPr внутри абзаца)."""
        mapping = []
        current = 0
        for child in doc.element.body.iterchildren():
            if child.tag == qn('w:p'):
                mapping.append(current)
                pPr = child.find(qn('w:pPr'))
                if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                    current += 1
        return mapping

    def _has_page_number_field(self, header) -> bool:
        xml_str = header._element.xml
        return ("instrText" in xml_str and "PAGE" in xml_str) or ("w:fldSimple" in xml_str and "PAGE" in xml_str)

    def _get_effective_font_size(self, run, paragraph) -> Optional[float]:
        if run.font.size:
            return run.font.size.pt
        if paragraph.style and paragraph.style.font and paragraph.style.font.size:
            return paragraph.style.font.size.pt
        return None

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "AppendixFormatRule"

        paragraphs = classify_paragraphs(doc)

        appendix_idx = None
        for idx, paragraph in enumerate(paragraphs):
            text = paragraph.text.strip()
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue
            if re.match(self.HEADER_REGEX, text.upper()):
                appendix_idx = idx
                break

        if appendix_idx is None:
            return violations  # Приложения в документе нет

        header_para = paragraphs[appendix_idx]
        header_text = header_para.text.strip()[:30]
        loc = f"Приложение (абзац {appendix_idx + 1}: '{header_text}')"

        # 1. Заголовок выравнивается по правому краю
        align = header_para.alignment
        if align is None and header_para.style and header_para.style.paragraph_format:
            align = header_para.style.paragraph_format.alignment
        if align != WD_ALIGN_PARAGRAPH.RIGHT:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.ERROR,
                    location=loc,
                    message="Заголовок приложения должен быть выровнен по правому краю."
                )
            )

        is_bold = any(run.bold for run in header_para.runs if run.text.strip())
        if not is_bold and header_para.style and header_para.style.font:
            is_bold = header_para.style.font.bold or False
        if not is_bold:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location=loc,
                    message="Заголовок приложения должен быть выделен полужирным шрифтом."
                )
            )

        # 2. В приложении шрифт 12 пт и межстрочный интервал 1
        for offset, paragraph in enumerate(paragraphs[appendix_idx + 1:], start=appendix_idx + 1):
            text = paragraph.text.strip()
            if not text:
                continue

            p_loc = f"Приложение (абзац {offset + 1}: '{text[:30]}...')"

            for r_idx, run in enumerate(paragraph.runs, start=1):
                if not run.text.strip():
                    continue
                font_size = self._get_effective_font_size(run, paragraph)
                if font_size and abs(font_size - 12.0) > 0.1:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"{p_loc}, фрагмент {r_idx}",
                            message=f"Размер шрифта в приложении {font_size} pt, а допустим 12 pt."
                        )
                    )

            line_sp = paragraph.paragraph_format.line_spacing
            if line_sp is None and paragraph.style and paragraph.style.paragraph_format:
                line_sp = paragraph.style.paragraph_format.line_spacing
            if line_sp is not None and abs(float(line_sp) - 1.0) > 0.05:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=p_loc,
                        message=f"Междустрочный интервал в приложении {line_sp}, а должен быть 1."
                    )
                )

        # 3. Страницы приложения не нумеруются
        section_indices = self._get_paragraph_section_indices(doc)
        if appendix_idx < len(section_indices):
            appendix_section = section_indices[appendix_idx]
            if appendix_section > 0:  # Если приложение в первой секции — отделить нумерацию нельзя
                for sec_idx, section in enumerate(doc.sections):
                    if sec_idx < appendix_section:
                        continue
                    if self._has_page_number_field(section.header):
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.ERROR,
                                location=f"Секция {sec_idx + 1} (Верхний колонтитул)",
                                message="В приложении страницы не должны нумероваться."
                            )
                        )

        return violations

class FiguresRule(BaseRule):
    CAPTION_REGEX = re.compile(r'^Рис(унок)?\.?\s*(\d+)', re.IGNORECASE)
    REF_REGEX = re.compile(r'\(\s*см\.\s*рис\.?\s*(\d+)\s*\)', re.IGNORECASE)
    PLAIN_REF_REGEX = re.compile(r'\(\s*рис\.?\s*(\d+)', re.IGNORECASE)

    def _figure_kind(self, paragraph) -> Optional[str]:
        """Тип абзаца с рисунком: 'inline', 'anchored' или None."""
        xml = paragraph._element.xml
        if '<w:drawing' not in xml and '<w:pict' not in xml:
            return None
        has_visual = ('<pic:pic' in xml or 'c:chart' in xml
                      or '<v:imagedata' in xml or '<w:object' in xml)
        if not has_visual:
            return None  # Надписи и фигуры без изображения рисунком не считаются
        if '<wp:anchor' in xml:
            return 'anchored'
        return 'inline'

    def _is_centered(self, paragraph) -> bool:
        align = paragraph.alignment
        if align is None and paragraph.style and paragraph.style.paragraph_format:
            align = paragraph.style.paragraph_format.alignment
        return align == WD_ALIGN_PARAGRAPH.CENTER

    def _has_red_line(self, paragraph) -> bool:
        indent = paragraph.paragraph_format.first_line_indent
        if indent is None and paragraph.style and paragraph.style.paragraph_format:
            indent = paragraph.style.paragraph_format.first_line_indent
        return indent is not None and indent.cm > 0.05

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "FiguresRule"

        paragraphs = classify_paragraphs(doc)

        caption_numbers = []
        referenced_numbers = []

        for idx, paragraph in enumerate(paragraphs):
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue

            kind = self._figure_kind(paragraph)
            if kind is None:
                continue

            loc = f"Рисунок в абзаце {idx + 1}"

            # 1. Рисунок выравнивается по центру (только встроенные рисунки)
            if kind == 'inline' and not self._is_centered(paragraph):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Рисунок должен быть выровнен по центру."
                    )
                )

            if kind == 'inline' and not self._has_red_line(paragraph):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="У рисунка должна быть красная строка."
                    )
                )

            # 2. Подпись под рисунком: 'Рис. N. Название'
            caption_para = None
            caption_text = None
            own_lines = paragraph.text.split('\n')
            if len(own_lines) > 1 and own_lines[-1].strip():
                caption_text = own_lines[-1].strip()
            else:
                for nxt in range(idx + 1, len(paragraphs)):
                    if paragraphs[nxt].text.strip():
                        caption_para = paragraphs[nxt]
                        caption_text = caption_para.text.strip()
                        break

            if caption_text and self.CAPTION_REGEX.match(caption_text):
                num = int(self.CAPTION_REGEX.match(caption_text).group(2))
                caption_numbers.append(num)

                cap_para = caption_para if caption_para is not None else paragraph
                if not self._is_centered(cap_para):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Подпись рисунка (абзац {idx + 1}: '{caption_text[:30]}')",
                            message="Подпись к рисунку должна быть выровнена по центру."
                        )
                    )
                if not self._has_red_line(cap_para):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"Подпись рисунка (абзац {idx + 1}: '{caption_text[:30]}')",
                            message="У подписи к рисунку должна быть красная строка."
                        )
                    )
            else:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="После рисунка отсутствует подпись формата 'Рис. N. Название'."
                    )
                )

        # 3. Ссылки на рисунки: (см.рис. N) в конце абзаца перед точкой
        for idx, paragraph in enumerate(paragraphs):
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue

            text = paragraph.text
            if not text.strip():
                continue

            for match in self.REF_REGEX.finditer(text):
                num = int(match.group(1))
                referenced_numbers.append(num)

                tail = text[match.end():].strip()
                if tail not in ("", "."):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"Абзац {idx + 1} ('{text.strip()[:30]}...')",
                            message=f"Ссылка на рисунок '(см.рис. {num})' должна стоять в конце абзаца перед точкой."
                        )
                    )

            # Ссылка без 'см.': '(рис. 5)' вместо '(см.рис. 5)'
            cleaned = self.REF_REGEX.sub('', text)
            for match in self.PLAIN_REF_REGEX.finditer(cleaned):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=f"Абзац {idx + 1} ('{text.strip()[:30]}...')",
                        message=f"Ссылка на рисунок должна иметь вид '(см.рис. {match.group(1)})'."
                    )
                )

        # 4. Соответствие ссылок и подписей
        for num in sorted(set(referenced_numbers)):
            if num not in caption_numbers:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location="Документ",
                        message=f"Ссылка на несуществующий рисунок {num} (подпись 'Рис. {num}.' не найдена)."
                    )
                )

        for num in set(caption_numbers):
            if num not in referenced_numbers:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location="Документ",
                        message=f"На рисунок {num} нет ссылки в тексте (должна быть: '(см.рис. {num})')."
                    )
                )

        # 5. Сквозная нумерация рисунков
        expected = list(range(1, len(caption_numbers) + 1))
        if sorted(caption_numbers) != expected and caption_numbers:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location="Документ",
                    message=f"Нумерация рисунков должна быть сквозной начиная с 1 (найдены: {sorted(caption_numbers)})."
                )
            )

        return violations

class TablesFormatRule(BaseRule):
    LABEL_REGEX = re.compile(r'^Таблица\s+(\d+)\s*$', re.IGNORECASE)
    CONTINUATION_REGEX = re.compile(r'^Продолжение таблицы\s+(\d+)\s*$', re.IGNORECASE)
    REF_REGEX = re.compile(r'\(\s*см\.\s*табл\.?\s*(\d+)\s*\)', re.IGNORECASE)

    def __init__(self, target_font_name: str = "Times New Roman", target_font_size: float = 14.0):
        self.font_name = target_font_name
        self.font_size = target_font_size

    @staticmethod
    def _body_sequence(paragraphs: List, tables: List):
        """Последовательность ('p', абзац, индекс) / ('tbl', таблица, индекс) в порядке следования."""
        items = []
        p_idx = 0
        t_idx = 0
        para_iter = iter(paragraphs)
        tbl_iter = iter(tables)
        if paragraphs:
            body = paragraphs[0]._element.getparent()
            for child in body.iterchildren():
                if child.tag == qn('w:p'):
                    para = next(para_iter, None)
                    if para is not None:
                        items.append(('p', para, p_idx))
                        p_idx += 1
                elif child.tag == qn('w:tbl'):
                    table = next(tbl_iter, None)
                    if table is not None:
                        items.append(('tbl', table, t_idx))
                        t_idx += 1
        return items

    def _get_effective_font_name(self, run, paragraph) -> Optional[str]:
        if run.font.name:
            return run.font.name
        if paragraph.style and paragraph.style.font and paragraph.style.font.name:
            return paragraph.style.font.name
        return None

    def _get_effective_font_size(self, run, paragraph) -> Optional[float]:
        if run.font.size:
            return run.font.size.pt
        if paragraph.style and paragraph.style.font and paragraph.style.font.size:
            return paragraph.style.font.size.pt
        return None

    def _is_centered(self, paragraph) -> bool:
        align = paragraph.alignment
        if align is None and paragraph.style and paragraph.style.paragraph_format:
            align = paragraph.style.paragraph_format.alignment
        return align == WD_ALIGN_PARAGRAPH.CENTER

    def _has_red_line(self, paragraph) -> bool:
        indent = paragraph.paragraph_format.first_line_indent
        if indent is None and paragraph.style and paragraph.style.paragraph_format:
            indent = paragraph.style.paragraph_format.first_line_indent
        return indent is not None and indent.cm > 0.05

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "TablesFormatRule"

        paragraphs = classify_paragraphs(doc)
        items = self._body_sequence(paragraphs, doc.tables)

        labeled_numbers = []
        referenced_numbers = []

        for pos, (kind, obj, obj_idx) in enumerate(items):
            if kind != 'tbl':
                continue

            table = obj
            table_loc = f"Таблица №{obj_idx + 1} по порядку в документе"

            # Контекст: до 3 непустых абзацев перед таблицей
            back = []
            j = pos - 1
            while j >= 0 and len(back) < 3:
                if items[j][0] == 'p' and items[j][1].text.strip():
                    back.append(items[j])
                j -= 1

            label_para = None
            title_para = None
            is_continuation = False

            if back:
                m = self.CONTINUATION_REGEX.match(back[0][1].text.strip())
                if m:
                    is_continuation = True
                    if back[0][1].alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.ERROR,
                                location=f"Продолжение таблицы {m.group(1)}",
                                message="Надпись 'Продолжение таблицы N' должна быть выровнена по правому краю."
                            )
                        )
                else:
                    m = self.LABEL_REGEX.match(back[0][1].text.strip())
                    if m:
                        label_para = back[0][1]
                    elif len(back) > 1:
                        m = self.LABEL_REGEX.match(back[1][1].text.strip())
                        if m:
                            label_para = back[1][1]
                            title_para = back[0][1]

            if is_continuation:
                continue  # Продолжение таблицы: надпись и название не требуются

            if label_para is None:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=table_loc,
                        message="Перед таблицей не найдена надпись 'Таблица N' (выравнивание по правому краю)."
                    )
                )
            else:
                num = int(self.LABEL_REGEX.match(label_para.text.strip()).group(1))
                labeled_numbers.append(num)

                if label_para.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Надпись 'Таблица {num}'",
                            message="Надпись 'Таблица N' должна быть выровнена по правому краю."
                        )
                    )

                if title_para is None:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.ERROR,
                            location=f"Таблица {num}",
                            message="После надписи 'Таблица N' отсутствует название таблицы."
                        )
                    )
                else:
                    title_text = title_para.text.strip()[:30]
                    if not self._is_centered(title_para):
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.ERROR,
                                location=f"Название таблицы {num} ('{title_text}')",
                                message="Название таблицы должно быть выровнено по центру."
                            )
                        )
                    if not self._has_red_line(title_para):
                        violations.append(
                            Violation(
                                rule_name=rule_name,
                                severity=Severity.WARNING,
                                location=f"Название таблицы {num} ('{title_text}')",
                                message="У названия таблицы должна быть красная строка."
                            )
                        )

            # Шрифт и интервал в ячейках таблицы
            bad_font = []
            bad_spacing = []
            for row in table.rows:
                for cell in row.cells:
                    for cell_para in cell.paragraphs:
                        if not cell_para.text.strip():
                            continue
                        line_sp = cell_para.paragraph_format.line_spacing
                        if line_sp is None and cell_para.style and cell_para.style.paragraph_format:
                            line_sp = cell_para.style.paragraph_format.line_spacing
                        if line_sp is not None and abs(float(line_sp) - 1.0) > 0.05:
                            bad_spacing.append(cell_para.text.strip()[:20])
                        for run in cell_para.runs:
                            if not run.text.strip():
                                continue
                            font_name = self._get_effective_font_name(run, cell_para)
                            if font_name and font_name != self.font_name:
                                bad_font.append(f"шрифт '{font_name}'")
                            font_size = self._get_effective_font_size(run, cell_para)
                            if font_size and abs(font_size - self.font_size) > 0.1:
                                bad_font.append(f"размер {font_size} pt")

            if bad_font:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=table_loc,
                        message=f"В таблице должен быть шрифт {self.font_name} {int(self.font_size)} pt "
                                f"(обнаружено: {', '.join(sorted(set(bad_font))[:3])})."
                    )
                )

            if bad_spacing:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=table_loc,
                        message=f"Междустрочный интервал в таблице должен быть 1 "
                                f"(обнаружен отличный в {len(bad_spacing)} фрагментах)."
                    )
                )

            # Названия столбцов (первая строка) — по центру по горизонтали и вертикали
            if table.rows:
                bad_header = []
                for cell in table.rows[0].cells:
                    cell_para = cell.paragraphs[0] if cell.paragraphs else None
                    if cell_para is None:
                        continue
                    align = cell_para.alignment
                    if align is None and cell_para.style and cell_para.style.paragraph_format:
                        align = cell_para.style.paragraph_format.alignment
                    if align != WD_ALIGN_PARAGRAPH.CENTER or cell.vertical_alignment != WD_ALIGN_VERTICAL.CENTER:
                        bad_header.append(cell.text.strip()[:20])
                if bad_header:
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=table_loc,
                            message=f"Названия столбцов должны быть выровнены по центру "
                                    f"(по горизонтали и вертикали); проблемы в: {', '.join(bad_header[:3])}."
                        )
                    )

        # Ссылки на таблицы: (см.табл. N) в конце абзаца перед точкой
        for p_idx, paragraph in enumerate(paragraphs, start=1):
            if getattr(paragraph, 'is_title_page', False) or getattr(paragraph, 'is_toc', False):
                continue
            text = paragraph.text
            if not text.strip():
                continue

            for match in self.REF_REGEX.finditer(text):
                num = int(match.group(1))
                referenced_numbers.append(num)
                tail = text[match.end():].strip()
                if tail not in ("", "."):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=f"Абзац {p_idx} ('{text.strip()[:30]}...')",
                            message=f"Ссылка на таблицу '(см.табл. {num})' должна стоять в конце абзаца перед точкой."
                        )
                    )

        for num in sorted(set(referenced_numbers)):
            if num not in labeled_numbers:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location="Документ",
                        message=f"Ссылка на несуществующую таблицу {num} (надпись 'Таблица {num}' не найдена)."
                    )
                )

        for num in set(labeled_numbers):
            if num not in referenced_numbers:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location="Документ",
                        message=f"На таблицу {num} нет ссылки в тексте (должна быть: '(см.табл. {num})')."
                    )
                )

        expected = list(range(1, len(labeled_numbers) + 1))
        if labeled_numbers and sorted(labeled_numbers) != expected:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location="Документ",
                    message=f"Нумерация таблиц должна быть сквозной начиная с 1 (найдены: {sorted(labeled_numbers)})."
                )
            )

        return violations

class FormulasRule(BaseRule):
    NUMBERING_REGEX = re.compile(r'\((\d+)\)\s*$')

    def _is_display_formula(self, paragraph) -> bool:
        """Выносная формула: oMathPara либо короткий абзац с oMath."""
        xml = paragraph._element.xml
        if '<m:oMath' not in xml:
            return False
        if '<m:oMathPara' in xml:
            return True
        text = self.NUMBERING_REGEX.sub('', paragraph.text).strip()
        return len(text) <= 25

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "FormulasRule"

        paragraphs = classify_paragraphs(doc)
        formula_indices = set()

        for idx, paragraph in enumerate(paragraphs):
            if not self._is_display_formula(paragraph):
                continue
            formula_indices.add(idx)

        formula_numbers = []

        for idx in sorted(formula_indices):
            paragraph = paragraphs[idx]
            text = paragraph.text.strip()
            loc = f"Формула (абзац {idx + 1}: '{text[:30]}')"

            # 1. Формула выравнивается по центру
            align = paragraph.alignment
            if align is None and paragraph.style and paragraph.style.paragraph_format:
                align = paragraph.style.paragraph_format.alignment
            if align != WD_ALIGN_PARAGRAPH.CENTER:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.ERROR,
                        location=loc,
                        message="Формула должна быть выровнена по центру."
                    )
                )

            # 2. Справа проставляется нумерация, например (1)
            match = self.NUMBERING_REGEX.search(text)
            if match:
                formula_numbers.append(int(match.group(1)))
            else:
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="Справа от формулы должна проставляться нумерация, например (1)."
                    )
                )

            # 3. Перед формулой описание, заканчивающееся двоеточием
            prev_idx = idx - 1
            while prev_idx >= 0 and (not paragraphs[prev_idx].text.strip() or prev_idx in formula_indices):
                prev_idx -= 1
            if prev_idx >= 0:
                prev_text = paragraphs[prev_idx].text.strip()
                if not prev_text.endswith(':'):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message="Перед формулой должно быть описание, заканчивающееся двоеточием."
                        )
                    )

            # 4. После формулы описываются примененные обозначения
            next_idx = idx + 1
            while next_idx < len(paragraphs) and (not paragraphs[next_idx].text.strip() or next_idx in formula_indices):
                next_idx += 1
            if next_idx < len(paragraphs):
                next_text = paragraphs[next_idx].text.strip().lower()
                if not next_text.startswith('где'):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message="После формулы должны быть описаны примененные обозначения (например: 'где: ...')."
                        )
                    )

        expected = list(range(1, len(formula_numbers) + 1))
        if formula_numbers and sorted(formula_numbers) != expected:
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location="Документ",
                    message=f"Нумерация формул должна быть сквозной начиная с 1 (найдены: {sorted(formula_numbers)})."
                )
            )

        return violations

class ListsFormatRule(BaseRule):
    DASH_ITEM_REGEX = re.compile(r'^(–|-|—)\s*(.+)$')
    BRACKET_ITEM_REGEX = re.compile(r'^(\d+)\)\s*(.+)$')
    DOT_ITEM_REGEX = re.compile(r'^(\d+)\.\s+(.+)$')

    INDENT_TOLERANCE_CM = 0.15

    @staticmethod
    def _item_kind(text: str):
        """Тип пункта списка: 'dot', 'bracket', 'dash' или None."""
        if ListsFormatRule.DOT_ITEM_REGEX.match(text):
            return 'dot'
        if ListsFormatRule.BRACKET_ITEM_REGEX.match(text):
            return 'bracket'
        if ListsFormatRule.DASH_ITEM_REGEX.match(text):
            return 'dash'
        return None

    def _get_effective_left_indent(self, paragraph):
        indent = paragraph.paragraph_format.left_indent
        if indent is None and paragraph.style and paragraph.style.paragraph_format:
            indent = paragraph.style.paragraph_format.left_indent
        return indent

    def _get_effective_first_indent(self, paragraph):
        indent = paragraph.paragraph_format.first_line_indent
        if indent is None and paragraph.style and paragraph.style.paragraph_format:
            indent = paragraph.style.paragraph_format.first_line_indent
        return indent

    def _check_item(self, paragraphs, idx, kind, is_last, violations):
        rule_name = "ListsFormatRule"
        paragraph = paragraphs[idx]
        text = paragraph.text.strip()
        loc = f"Абзац {idx + 1} ('{text[:30]}...')"

        if kind == 'dot':
            match = self.DOT_ITEM_REGEX.match(text)
            number = int(match.group(1))
            body = match.group(2)
            expected_indent = 2.0 if number >= 10 else 1.75
        elif kind == 'bracket':
            match = self.BRACKET_ITEM_REGEX.match(text)
            body = match.group(2)
            expected_indent = 1.75
        else:
            body = self.DASH_ITEM_REGEX.match(text).group(2)
            expected_indent = 1.75

        # 1. Регистр первой буквы после маркера
        first_char = body[:1]
        if first_char.isalpha():
            if kind in ('dash', 'bracket') and first_char.isupper():
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"После {'тире' if kind == 'dash' else 'скобки'} список пишется с маленькой буквы."
                    )
                )
            if kind == 'dot' and first_char.islower():
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="После точки список пишется с большой буквы."
                    )
                )

        # 2. Знак в конце пункта
        last_char = text[-1:]
        if kind == 'dot':
            if last_char not in ('.',):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="В нумерованном списке с точкой в конце пункта всегда ставится точка."
                    )
                )
        elif is_last:
            if last_char not in ('.', ':'):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message="В последнем пункте списка ставится точка."
                    )
                )
        else:
            if last_char not in (';',):
                violations.append(
                    Violation(
                        rule_name=rule_name,
                        severity=Severity.WARNING,
                        location=loc,
                        message=f"Пункты {'маркированного' if kind == 'dash' else 'нумерованного'} списка "
                                f"заканчиваются точкой с запятой (кроме последнего)."
                    )
                )

        # 3. Отступ текста списка
        left = self._get_effective_left_indent(paragraph)
        first = self._get_effective_first_indent(paragraph)
        left_cm = left.cm if left is not None else None
        first_cm = first.cm if first is not None else None

        indent_values = [v for v in (left_cm, first_cm) if v is not None]
        if left_cm is not None and first_cm is not None:
            indent_values.append(left_cm + first_cm)
        if not any(abs(v - expected_indent) <= self.INDENT_TOLERANCE_CM for v in indent_values):
            violations.append(
                Violation(
                    rule_name=rule_name,
                    severity=Severity.WARNING,
                    location=loc,
                    message=f"Отступ текста списка должен составлять {expected_indent} см "
                            f"(фактически: левый {round(left_cm, 2) if left_cm is not None else '—'} см, "
                            f"красная строка {round(first_cm, 2) if first_cm is not None else '—'} см)."
                )
            )

    def check(self, doc: Document) -> List[Violation]:
        violations: List[Violation] = []
        rule_name = "ListsFormatRule"

        paragraphs = classify_paragraphs(doc)

        def is_centered(p):
            align = p.alignment
            if align is None and p.style and p.style.paragraph_format:
                align = p.style.paragraph_format.alignment
            return align == WD_ALIGN_PARAGRAPH.CENTER

        item_indices = set()
        for idx, p in enumerate(paragraphs):
            if getattr(p, 'is_title_page', False) or getattr(p, 'is_toc', False):
                continue
            text = p.text.strip()
            if not text:
                continue
            # Список литературы и приложения не проверяются как списки
            if re.match(r'^(СПИСОК\s+(ИСТОЧНИКОВ|ЛИТЕРАТУРЫ|ИСПОЛЬЗОВАННЫХ\s+ИСТОЧНИКОВ)|ПРИЛОЖЕНИЕ)', text.upper()):
                break
            if is_centered(p):
                continue
            if self._item_kind(text):
                item_indices.add(idx)

        for idx in sorted(item_indices):
            paragraph = paragraphs[idx]
            text = paragraph.text.strip()
            kind = self._item_kind(text)
            loc = f"Абзац {idx + 1} ('{text[:30]}...')"

            # Соседний предыдущий/следующий пункт того же списка (пустые строки не считаются)
            prev_item_idx = idx - 1
            while prev_item_idx >= 0 and not paragraphs[prev_item_idx].text.strip():
                prev_item_idx -= 1
            next_item_idx = idx + 1
            while next_item_idx < len(paragraphs) and not paragraphs[next_item_idx].text.strip():
                next_item_idx += 1

            has_prev_same = prev_item_idx in item_indices and self._item_kind(paragraphs[prev_item_idx].text.strip()) == kind
            has_next_same = next_item_idx in item_indices and self._item_kind(paragraphs[next_item_idx].text.strip()) == kind
            is_last = not has_next_same

            self._check_item(paragraphs, idx, kind, is_last, violations)

            # 4. Двоеточие в абзаце перед первым пунктом списка (из двух и более пунктов)
            if not has_prev_same and has_next_same and prev_item_idx >= 0:
                prev_text = paragraphs[prev_item_idx].text.strip()
                if not prev_text.endswith(':'):
                    violations.append(
                        Violation(
                            rule_name=rule_name,
                            severity=Severity.WARNING,
                            location=loc,
                            message="В абзаце перед списком ставится двоеточие."
                        )
                    )

        return violations
